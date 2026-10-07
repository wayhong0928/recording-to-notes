"""產生中性的 Word 公版（docxtpl 範本）。

用法（在 app\\ 底下）：uv run python templates/word/build_word_template.py
輸出：同資料夾的 中性範例.docx

範本吃的資料（由 minutes/render.py 從 md 拆出來）：
    project, meeting, location, time, attendees, recorder   字串，可空白
    contents     會議內容，字串清單
    decisions    決議事項，字串清單
    discussions  討論事項，[{title, items: [{no, ask, follow}]}]
                 ask 是詢問那行（含「XX詢問：」），follow 是底下縮排的回覆、補充、結論，字串清單
                 稱呼和動詞都照 md 原文，範本不改字；no 整份連續編號
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "微軟正黑體"
OUT = Path(__file__).parent / "中性範例.docx"

# A4 寬 21 cm，左右邊界各 2 cm，表格寬 17 cm
COLS = [Cm(2.8), Cm(14.2)]
LABEL_FILL = "D9D9D9"
SECTION_FILL = "BFBFBF"


def set_font(run, size=11, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def write(cell, lines, bold=False, align=None, size=11, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER):
    """把 cell 內容換成 lines，每行一個段落。"""
    first = cell.paragraphs[0]
    for i, text in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if align is not None:
            p.alignment = align
        if text:
            set_font(p.add_run(text), size=size, bold=bold)
    cell.vertical_alignment = valign


def label(cell, text):
    write(cell, [text], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(cell, LABEL_FILL)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("會議紀錄"), size=18, bold=True)

    table = doc.add_table(rows=12, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, COLS):
            cell.width = w
    # add_table 預設等分欄寬，表格層級的 gridCol 也要改，不然 LibreOffice 會照等分顯示
    for col, w in zip(table._tbl.tblGrid.gridCol_lst, COLS):
        col.w = w

    # 基本資料：每列一個欄位，標籤都在左邊
    basics = [("專案名稱", "project"), ("會議名稱", "meeting"), ("會議地點", "location"),
              ("會議時間", "time"), ("與會人員", "attendees"), ("記錄人", "recorder")]
    for r, (name, key) in enumerate(basics):
        cells = table.rows[r].cells
        label(cells[0], name)
        write(cells[1], ["{{ %s }}" % key])

    # 三個段落：標題列＋內容列，都橫跨兩欄
    sections = [
        (6, "會議內容", ["{%p for x in contents %}", "{{ loop.index }}. {{ x }}", "{%p endfor %}"]),
        (8, "決議事項", ["{%p for x in decisions %}", "{{ loop.index }}. {{ x }}", "{%p endfor %}"]),
        (10, "討論事項", [
            "{%p for d in discussions %}",
            # 議題之間空一行，第一個議題直接接在標題列底下
            "{%p if not loop.first %}",
            "",
            "{%p endif %}",
            "{%p if d.title %}",
            "{{ d.title }}",
            "{%p endif %}",
            "{%p for q in d['items'] %}",
            "{{ q.no }}. {{ q.ask }}",
            "{%p for line in q.follow %}",
            "{{ line }}",
            "{%p endfor %}",
            "{%p endfor %}",
            "{%p endfor %}",
        ]),
    ]
    for r, name, body in sections:
        head = table.rows[r].cells
        head_cell = head[0].merge(head[1])
        write(head_cell, [name], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(head_cell, SECTION_FILL)
        content = table.rows[r + 1].cells
        body_cell = content[0].merge(content[1])
        write(body_cell, body, valign=WD_CELL_VERTICAL_ALIGNMENT.TOP)
        # 回覆、補充、結論縮排在詢問底下
        for p in body_cell.paragraphs:
            if p.text == "{{ line }}":
                p.paragraph_format.left_indent = Cm(0.75)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"done: {OUT}")


if __name__ == "__main__":
    main()
