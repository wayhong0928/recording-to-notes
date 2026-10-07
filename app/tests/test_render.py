import json
import shutil
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document

from minutes import paths, render
from minutes.cli import main

EXAMPLES = paths.templates() / "examples"
EXAMPLE = "雙十一檔期準備會議"


def out_json(capsys):
    return json.loads(capsys.readouterr().out)


def table_shape(docx: Path) -> list:
    """表格每列、每格（合併的格只算一次）的段落文字和縮排。docx 帶時間戳，只能比內容。"""
    rows = []
    for row in Document(docx).tables[0].rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append([(p.text, bool(p.paragraph_format.left_indent)) for p in cell.paragraphs])
        rows.append(cells)
    return rows


@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "2026-10-01_雙十一檔期準備會議"
    f.mkdir()
    shutil.copy(EXAMPLES / f"{EXAMPLE}.md", f / render.MINUTES_FILE)
    return f


# ---- 解析規則 ----

def test_fields_full_and_half_width_colon():
    data = render.parse("專案名稱：甲\n**會議名稱**: 乙\n- 會議地點 ： 丙\n記錄人：\n")
    assert (data["project"], data["meeting"], data["location"], data["recorder"]) == ("甲", "乙", "丙", "")


def test_numbering_gaps_and_continuation_lines():
    data = render.parse("## 決議事項\n1. 第一項\n5. 第二項\n接在第二項後面\n- 第三項\n")
    assert data["decisions"] == ["第一項", "第二項接在第二項後面", "第三項"]


def test_decimal_is_not_a_number():
    data = render.parse("## 決議事項\n3.5 倍的預算\n")
    assert data["decisions"] == ["3.5 倍的預算"]


def test_discussion_structure_and_renumbering():
    md = ("## 討論事項\n議題甲\n7. 客戶詢問：問一\n   我方回覆：答一\n\t補充\n"
          "議題乙\n9. 客戶詢問：問二\n　結論：全形縮排\n- 後續：用減號\n")
    topics = render.parse(md)["discussions"]
    assert [t["title"] for t in topics] == ["議題甲", "議題乙"]
    assert topics[0]["items"][0] == {"ask": "客戶詢問：問一", "follow": ["我方回覆：答一", "補充"], "no": 1}
    assert topics[1]["items"][0] == {"ask": "客戶詢問：問二", "follow": ["結論：全形縮排", "後續：用減號"], "no": 2}


def test_reply_without_indent_becomes_topic():
    topics = render.parse("## 討論事項\n1. 客戶詢問：問\n我方回覆：沒縮排\n")["discussions"]
    assert [t["title"] for t in topics] == ["", "我方回覆：沒縮排"]


def test_special_characters_survive(folder):
    (folder / render.MINUTES_FILE).write_text("會議名稱：A&B <測試>\n\n## 決議事項\n1. 成本 < 預算 & 3.5 倍\n",
                                             encoding="utf-8")
    out = Path(render.render(folder, open_word=False)["docx"])
    text = "\n".join(p.text for row in Document(out).tables[0].rows for c in row.cells for p in c.paragraphs)
    assert "A&B <測試>" in text
    assert "1. 成本 < 預算 & 3.5 倍" in text


# ---- 產生 ----

def test_example_matches_reference_docx(folder):
    out = Path(render.render(folder, open_word=False)["docx"])
    assert table_shape(out) == table_shape(EXAMPLES / f"{EXAMPLE}.docx")


def test_output_name_has_time_and_never_overwrites(folder):
    now = datetime(2026, 10, 1, 17, 5, 12)
    first = render.output_path(folder, now)
    assert first.name == "2_會議紀錄_20261001170512.docx"
    first.write_bytes(b"")
    second = render.output_path(folder, now)
    assert second.name == "2_會議紀錄_20261001170512_2.docx"
    second.write_bytes(b"")
    assert render.output_path(folder, now).name == "2_會議紀錄_20261001170512_3.docx"


def test_rerender_keeps_old_files(folder):
    a = render.render(folder, open_word=False)["docx"]
    b = render.render(folder, open_word=False)["docx"]
    assert a != b and Path(a).exists() and Path(b).exists()


def test_missing_md(tmp_path):
    with pytest.raises(render.RenderError, match="找不到 2_會議紀錄.md"):
        render.render(tmp_path, open_word=False)


def test_open_failure_is_not_an_error(folder, monkeypatch):
    def boom(*a, **k):
        raise OSError("no app")
    monkeypatch.setattr(render.os, "startfile", boom, raising=False)
    monkeypatch.setattr(render.subprocess, "run", boom)
    result = render.render(folder, open_word=True)
    assert result["ok"] and not result["opened"]
    assert "沒辦法自動打開 Word" in result["open_note"]


# ---- 指紋 ----

def test_no_warning_when_untouched(folder):
    render.render(folder, open_word=False)
    second = render.render(folder, open_word=False)
    assert second["previous_edited"] is False
    assert render.status(folder)["edited"] is False


def test_warning_when_edited_but_still_renders(folder):
    first = Path(render.render(folder, open_word=False)["docx"])
    with first.open("ab") as f:
        f.write(b"edited")
    st = render.status(folder)
    assert st["edited"] and st["latest"] == first.name and st["modified"]
    second = render.render(folder, open_word=False)
    assert second["previous_edited"] is True and Path(second["docx"]).exists()
    assert "提醒：上一版" in render.render_text(second)
    assert render.status(folder)["edited"] is False  # 最新的是剛產生那份


def test_status_ignores_deleted_latest(folder):
    first = Path(render.render(folder, open_word=False)["docx"])
    second = Path(render.render(folder, open_word=False)["docx"])
    second.unlink()
    assert render.status(folder)["latest"] == first.name


# ---- 指令 ----

def test_word_status_cli(folder, capsys):
    assert main(["word-status", str(folder)]) == 0
    assert "還沒有產生過 Word" in capsys.readouterr().out
    assert main(["render", str(folder), "--no-open"]) == 0
    assert "已產生 2_會議紀錄_" in capsys.readouterr().out
    assert main(["word-status", str(folder), "--json"]) == 0
    data = out_json(capsys)
    assert data["ok"] is True and data["edited"] is False and data["latest"].startswith("2_會議紀錄_")
    assert main(["word-status", str(folder)]) == 0
    assert "之後沒有被改過" in capsys.readouterr().out


def test_render_cli_json(folder, capsys):
    assert main(["--json", "render", str(folder), "--no-open"]) == 0
    data = out_json(capsys)
    assert data["ok"] is True and data["counts"] == {"contents": 3, "decisions": 3, "topics": 2, "questions": 3}


def test_cli_errors_have_ok_and_error(tmp_path, capsys):
    assert main(["render", str(tmp_path), "--no-open", "--json"]) == 1
    data = out_json(capsys)
    assert data["ok"] is False and "找不到 2_會議紀錄.md" in data["error"]
    assert main(["word-status", str(tmp_path / "沒有這個"), "--json"]) == 1
    assert out_json(capsys)["ok"] is False
